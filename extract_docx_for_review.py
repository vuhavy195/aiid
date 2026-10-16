from pathlib import Path
import json
from docx import Document
from docx.oxml.ns import qn


src = Path(r"H:\AIID_DTA_GPR\Draft_FullPaper_CIEMB.docx")
out = Path(r"H:\AIID_DTA_GPR\.review")
out.mkdir(exist_ok=True)

doc = Document(src)
records = []

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        records.append({
            "kind": "paragraph",
            "index": i,
            "style": p.style.name if p.style else "",
            "text": text,
        })

tables = []
for ti, table in enumerate(doc.tables):
    rows = []
    for row in table.rows:
        rows.append([" ".join(cell.text.split()) for cell in row.cells])
    tables.append({"table": ti, "rows": rows})

headings = [r for r in records if r["style"].lower().startswith("heading")]

full_xml_text = []
for node in doc.element.body.iter():
    if node.tag == qn("w:t") and node.text:
        full_xml_text.append(node.text)

(out / "paragraphs.json").write_text(
    json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8"
)
(out / "tables.json").write_text(
    json.dumps(tables, ensure_ascii=False, indent=2), encoding="utf-8"
)
(out / "headings.json").write_text(
    json.dumps(headings, ensure_ascii=False, indent=2), encoding="utf-8"
)
(out / "all_body_text.txt").write_text("\n".join(full_xml_text), encoding="utf-8")

print(json.dumps({
    "paragraphs": len(records),
    "tables": len(tables),
    "headings": len(headings),
    "sections": len(doc.sections),
    "inline_shapes": len(doc.inline_shapes),
}, ensure_ascii=False))
